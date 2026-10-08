import io
import json
import os
import unittest
from datetime import date, datetime, timezone
from decimal import Decimal
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook

import app
import excel_export
from config import BIGQUERY_DATASET_ID, BIGQUERY_PROJECT_ID, BIGQUERY_TABLE_ID
from excel_export import (
    ExportTokenError,
    ExportTooLargeError,
    decrypt_token,
    encrypt_token,
    is_excel_export_request,
    write_workbook,
)
from sql_guardrails import validate_sql

SQL = "SELECT sport, COUNT(*) AS registration_count FROM `p.d.t` GROUP BY sport"


def openpyxl_temp_files():
    from openpyxl.worksheet._writer import ALL_TEMP_FILES

    return set(ALL_TEMP_FILES)


def parse_events(text):
    events = []
    for block in text.strip().split("\n\n"):
        lines = block.split("\n")
        events.append((lines[0][len("event: "):], json.loads(lines[1][len("data: "):])))
    return events


def sheet_rows(path_or_bytes):
    workbook = load_workbook(
        io.BytesIO(path_or_bytes) if isinstance(path_or_bytes, bytes) else path_or_bytes
    )
    return workbook, [list(row) for sheet in workbook.worksheets for row in sheet.iter_rows(values_only=True)]


class ExportIntentTests(unittest.TestCase):
    def test_detects_export_requests(self):
        for question in (
            "Export this data to Excel.",
            "Can you download these results in Excel?",
            "Give me the above results in an Excel sheet.",
            "Generate an Excel report.",
            "Export the complete data.",
            "Can you provide this analysis in XLSX format?",
            "Show athlete registrations grouped by sport and gender, and export the results to Excel.",
            "هل يمكنك تصدير هذه البيانات إلى ملف إكسل؟",
            "حمّل النتائج بصيغة Excel",
        ):
            with self.subTest(question=question):
                self.assertTrue(is_excel_export_request(question))

    def test_ignores_normal_questions(self):
        for question in (
            "How many registrations are there for each sport?",
            "Which sports have the most active athletes?",
            "hello",
            "كم عدد التسجيلات لكل رياضة؟",
        ):
            with self.subTest(question=question):
                self.assertFalse(is_excel_export_request(question))


class ExportTokenTests(unittest.TestCase):
    def test_round_trip_and_type_check(self):
        token = encrypt_token({"kind": "context", "sql": SQL})
        self.assertEqual(decrypt_token(token, 60, "context")["sql"], SQL)
        with self.assertRaises(ExportTokenError):
            decrypt_token(token, 60, "download")

    def test_tampered_and_expired_tokens_are_rejected(self):
        token = encrypt_token({"kind": "download", "sql": SQL})
        with self.assertRaises(ExportTokenError):
            decrypt_token(token[:-4] + "AAAA", 60, "download")
        with self.assertRaises(ExportTokenError):
            decrypt_token(token, -1, "download")

    def test_token_does_not_expose_sql(self):
        self.assertNotIn("SELECT", encrypt_token({"kind": "context", "sql": SQL}))


class WorkbookTests(unittest.TestCase):
    def write(self, headers, rows, **kwargs):
        path, count = write_workbook(headers, rows, **kwargs)
        self.addCleanup(lambda: os.path.exists(path) and os.remove(path))
        return path, count

    def test_values_types_and_arabic_are_preserved(self):
        rows = [
            ("كرة القدم", 32795, 1.5, Decimal("2.25"), date(2024, 1, 31), True, None),
            ("Football", 10**16, 0.1, Decimal("3"), date(2024, 2, 1), False, "x"),
        ]
        path, count = self.write(
            ["الرياضة", "Registrations", "Ratio", "Amount", "Day", "Active", "Note"],
            rows,
            language="ar",
        )

        workbook, loaded = sheet_rows(path)
        sheet = workbook.active
        self.assertEqual(count, 2)
        self.assertTrue(sheet.sheet_view.rightToLeft)
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(loaded[0][0], "الرياضة")
        self.assertEqual(loaded[1][:2], ["كرة القدم", 32795])
        self.assertEqual(loaded[1][2], 1.5)
        self.assertEqual(loaded[1][4].date(), date(2024, 1, 31))
        self.assertIs(loaded[1][5], True)
        self.assertEqual(loaded[2][1], str(10**16))  # beyond Excel's exact-integer range, kept exact as text
        self.assertEqual(loaded[2][0], "Football")

    def test_timezone_aware_datetimes_are_stored_as_utc(self):
        path, _ = self.write(
            ["At"],
            [(datetime(2024, 5, 1, 12, 0, tzinfo=timezone.utc),)],
        )
        self.assertEqual(sheet_rows(path)[1][1][0], datetime(2024, 5, 1, 12, 0))

    def test_formula_like_text_is_never_evaluated(self):
        path, _ = self.write(["Name"], [("=HYPERLINK(\"http://x\",\"y\")",), ("#N/A",)])
        workbook, loaded = sheet_rows(path)
        cells = [row[0] for row in workbook.active.iter_rows(min_row=2)]
        self.assertEqual([cell.data_type for cell in cells], ["s", "s"])
        self.assertEqual(loaded[1][0], "=HYPERLINK(\"http://x\",\"y\")")

    def test_long_exports_are_split_across_sheets(self):
        path, count = self.write(["n"], [(index,) for index in range(5)], sheet_max_rows=2)
        workbook, loaded = sheet_rows(path)
        self.assertEqual(count, 5)
        self.assertEqual(workbook.sheetnames, ["Data", "Data 2", "Data 3"])
        self.assertEqual([row[0] for row in loaded if isinstance(row[0], int)], [0, 1, 2, 3, 4])

    def test_row_limit_is_enforced_and_temp_file_removed(self):
        before = openpyxl_temp_files()
        created = []
        real_mkstemp = excel_export.tempfile.mkstemp

        def tracking_mkstemp(*args, **kwargs):
            handle, path = real_mkstemp(*args, **kwargs)
            created.append(path)
            return handle, path

        with patch("excel_export.tempfile.mkstemp", tracking_mkstemp):
            with self.assertRaises(ExportTooLargeError):
                write_workbook(["n"], iter([(1,), (2,), (3,)]), max_rows=2)

        self.assertTrue(created)
        self.assertFalse(os.path.exists(created[0]))
        self.assertEqual(openpyxl_temp_files(), before)

    def test_empty_result_still_has_headers_and_duplicates_are_made_unique(self):
        path, count = self.write(["Sport", "Sport"], [])
        self.assertEqual(count, 0)
        self.assertEqual(sheet_rows(path)[1], [["Sport", "Sport (2)"]])

    def test_friendly_headers(self):
        self.assertEqual(excel_export.friendly_header("registration_count"), "Registration Count")
        self.assertEqual(excel_export.friendly_header("Sport"), "Sport")


class ExportValidationTests(unittest.TestCase):
    query = f"SELECT sport FROM `{BIGQUERY_PROJECT_ID}.{BIGQUERY_DATASET_ID}.{BIGQUERY_TABLE_ID}`"

    def test_default_validation_still_caps_rows_but_export_mode_does_not(self):
        with patch("sql_guardrails.get_allowed_columns", return_value={"sport"}):
            self.assertTrue(validate_sql(self.query).endswith("LIMIT 100"))
            self.assertNotIn("LIMIT", validate_sql(self.query, row_limit=None))
            self.assertIn("LIMIT 5", validate_sql(self.query + " LIMIT 5", row_limit=None))

    def test_export_mode_keeps_every_other_guardrail(self):
        with patch("sql_guardrails.get_allowed_columns", return_value={"sport"}):
            with self.assertRaises(ValueError):
                validate_sql("DELETE FROM `a.b.c` WHERE 1=1", row_limit=None)
            with self.assertRaises(ValueError):
                validate_sql(self.query.replace("sport", "secret"), row_limit=None)


class ExportChatFlowTests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app.app)

    def ask(self, question, *, language="en", export_context=None, **patches):
        defaults = {
            "app.extract_export_analysis_question": None,
            "app.get_export_row_count": 3,
        }
        defaults.update(patches)
        stack = [
            patch(name, side_effect=value)
            if isinstance(value, BaseException) or callable(value)
            else patch(name, return_value=value)
            for name, value in defaults.items()
        ]
        stack.append(patch("app.validate_sql", side_effect=lambda sql, row_limit=100: sql))
        stack.append(patch("app.generate_sql", return_value=SQL))
        mocks = [context.start() for context in stack]
        for context in stack:
            self.addCleanup(context.stop)

        response = self.client.post(
            "/ask/stream",
            json={"question": question, "language": language, "export_context": export_context},
        )
        self.assertEqual(response.status_code, 200)
        return parse_events(response.text), {"generate_sql": mocks[-1], "validate_sql": mocks[-2]}

    def tokens(self, events):
        return "".join(payload["text"] for name, payload in events if name == "token")

    def test_data_answer_emits_context_for_follow_up_exports(self):
        with (
            patch("app.generate_sql", return_value=SQL),
            patch("app.validate_sql", side_effect=lambda sql, row_limit=100: sql),
            patch("app.execute_sql", return_value=(["sport"], [("Football",)])),
            patch("app.generate_answer_stream", return_value=iter(["Counts"])),
        ):
            response = self.client.post("/ask/stream", json={"question": "How many registrations are there for each sport?"})

        context = dict(parse_events(response.text))["context"]["export_context"]
        self.assertEqual(decrypt_token(context, 60, "context")["sql"], SQL)
        self.assertNotIn("SELECT", context)

    def test_follow_up_export_reuses_previous_validated_query_without_new_sql(self):
        context = app._export_context_token(SQL, "How many registrations are there for each sport?")
        events, mocks = self.ask("Export these results to Excel.", export_context=context)

        names = [name for name, _ in events]
        self.assertIn("export", names)
        self.assertEqual(self.tokens(events), "Your Excel report is ready.")
        mocks["generate_sql"].assert_not_called()
        mocks["validate_sql"].assert_called_once_with(SQL, row_limit=None)

        download = dict(events)["export"]
        self.assertEqual(download["row_count"], 3)
        self.assertTrue(download["download_url"].startswith("/exports/"))
        self.assertNotIn("SELECT", download["download_url"])

    def test_direct_export_runs_the_new_query_in_the_same_request(self):
        events, mocks = self.ask(
            "Show athlete registrations grouped by sport and gender, and export the results to Excel.",
            **{"app.extract_export_analysis_question": "Show athlete registrations grouped by sport and gender"},
        )

        mocks["generate_sql"].assert_called_once_with(
            "Show athlete registrations grouped by sport and gender", "en"
        )
        self.assertEqual(self.tokens(events), "Your Excel report is ready.")
        self.assertIn("export", dict(events))
        context = dict(events)["context"]["export_context"]
        self.assertEqual(decrypt_token(context, 60, "context")["sql"], SQL)

    def test_export_without_previous_results_asks_which_data(self):
        events, mocks = self.ask("Export this data to Excel.")

        self.assertIn("Which NSP data would you like to export", self.tokens(events))
        self.assertNotIn("export", dict(events))
        mocks["generate_sql"].assert_not_called()

    def test_export_is_not_treated_as_out_of_scope(self):
        events, _ = self.ask("Generate an Excel report.")
        self.assertNotEqual(self.tokens(events), app.OUT_OF_SCOPE_CLARIFICATION)

    def test_other_conversation_context_is_ignored_when_invalid_or_expired(self):
        for context in ("not-a-token", encrypt_token({"kind": "download", "sql": SQL})):
            with self.subTest(context=context[:12]):
                events, mocks = self.ask("Export these results to Excel.", export_context=context)
                self.assertIn("Which NSP data would you like to export", self.tokens(events))
                mocks["validate_sql"].assert_not_called()

    def test_no_matching_records(self):
        events, _ = self.ask(
            "Export these results to Excel.",
            export_context=app._export_context_token(SQL, "q"),
            **{"app.get_export_row_count": 0},
        )
        self.assertEqual(self.tokens(events), "No matching data was found, so no Excel file was created.")
        self.assertNotIn("export", dict(events))

    def test_large_exports_over_the_limit_are_refused(self):
        with patch("app.EXPORT_MAX_ROWS", 10):
            events, _ = self.ask(
                "Export these results to Excel.",
                export_context=app._export_context_token(SQL, "q"),
                **{"app.get_export_row_count": 11},
            )
        self.assertIn("11 rows", self.tokens(events))
        self.assertIn("10-row limit", self.tokens(events))
        self.assertNotIn("export", dict(events))

    def test_query_over_byte_limit_gets_a_clear_message(self):
        events, _ = self.ask(
            "Export these results to Excel.",
            export_context=app._export_context_token(SQL, "q"),
            **{"app.get_export_row_count": ValueError("too many bytes")},
        )
        self.assertIn("too large to process", self.tokens(events))

    def test_unexpected_failure_shows_a_meaningful_message(self):
        def boom(*_):
            raise RuntimeError("bigquery down")

        events, _ = self.ask(
            "Export these results to Excel.",
            export_context=app._export_context_token(SQL, "q"),
            **{"app.get_export_row_count": boom},
        )
        self.assertEqual(self.tokens(events), "I couldn't generate the Excel file. Please try again.")
        self.assertNotIn("bigquery down", self.tokens(events))

    def test_arabic_conversation_gets_arabic_messages(self):
        events, _ = self.ask(
            "صدّر هذه النتائج إلى Excel",
            language="ar",
            export_context=app._export_context_token(SQL, "q"),
        )
        self.assertEqual(self.tokens(events), "تقرير Excel جاهز.")

        events, _ = self.ask("صدّر هذه النتائج إلى Excel", language="ar")
        self.assertIn("البيانات", self.tokens(events))

    def test_metadata_requests_cannot_be_exported(self):
        events, mocks = self.ask(
            "Export the column names to Excel.",
            **{"app.extract_export_analysis_question": "What columns are available?"},
        )
        self.assertEqual(self.tokens(events), app.OUT_OF_SCOPE_CLARIFICATION)
        mocks["generate_sql"].assert_not_called()


class ExportDownloadTests(unittest.TestCase):
    columns = ["sport", "registration_count"]
    rows = [("Football", 32795), ("كرة اليد", 6107), ("Swimming", 5134)]

    def setUp(self):
        self.client = TestClient(app.app)

    def download(self, token, **patches):
        with (
            patch("app.validate_sql", side_effect=lambda sql, row_limit=100: sql),
            patch("app.stream_query_rows", return_value=(self.columns, iter(self.rows))),
            patch("app.translate_texts", return_value=["الرياضة", "عدد التسجيلات"]),
        ):
            return self.client.get(f"/exports/{token}")

    def test_downloaded_workbook_matches_bigquery_rows(self):
        token = encrypt_token({"kind": "download", "sql": SQL, "language": "en"})
        response = self.download(token)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["content-type"], excel_export.XLSX_MEDIA_TYPE)
        self.assertIn("attachment", response.headers["content-disposition"])
        self.assertIn(".xlsx", response.headers["content-disposition"])
        self.assertEqual(response.headers["cache-control"], "no-store")

        workbook, loaded = sheet_rows(response.content)
        self.assertEqual(loaded[0], ["Sport", "Registration Count"])
        self.assertEqual(loaded[1:], [list(row) for row in self.rows])
        self.assertFalse(workbook.active.sheet_view.rightToLeft)

    def test_arabic_workbook_has_arabic_headers_and_rtl(self):
        token = encrypt_token({"kind": "download", "sql": SQL, "language": "ar"})
        response = self.download(token)

        workbook, loaded = sheet_rows(response.content)
        self.assertEqual(loaded[0], ["الرياضة", "عدد التسجيلات"])
        self.assertEqual(loaded[1][0], "Football")  # database value is not translated
        self.assertTrue(workbook.active.sheet_view.rightToLeft)

    def test_arabic_header_translation_failure_falls_back_to_english(self):
        token = encrypt_token({"kind": "download", "sql": SQL, "language": "ar"})
        with (
            patch("app.validate_sql", side_effect=lambda sql, row_limit=100: sql),
            patch("app.stream_query_rows", return_value=(self.columns, iter(self.rows))),
            patch("app.translate_texts", side_effect=RuntimeError("unavailable")),
        ):
            response = self.client.get(f"/exports/{token}")

        self.assertEqual(sheet_rows(response.content)[1][0], ["Sport", "Registration Count"])

    def test_download_uses_the_export_mode_validator(self):
        token = encrypt_token({"kind": "download", "sql": SQL, "language": "en"})
        with (
            patch("app.validate_sql", side_effect=lambda sql, row_limit=100: sql) as validate,
            patch("app.stream_query_rows", return_value=(self.columns, iter(self.rows))),
        ):
            self.client.get(f"/exports/{token}")
        validate.assert_called_once_with(SQL, row_limit=None)

    def test_invalid_expired_or_wrong_type_links_are_rejected(self):
        for token in ("garbage", encrypt_token({"kind": "context", "sql": SQL})):
            with self.subTest(token=token[:10]):
                response = self.download(token)
                self.assertEqual(response.status_code, 410)

        token = encrypt_token({"kind": "download", "sql": SQL, "language": "en"})
        with patch("app.EXPORT_DOWNLOAD_TTL_SECONDS", -1):
            self.assertEqual(self.download(token).status_code, 410)

    def test_oversized_download_returns_413(self):
        token = encrypt_token({"kind": "download", "sql": SQL, "language": "en"})
        with patch("app.EXPORT_MAX_ROWS", 1):
            self.assertEqual(self.download(token).status_code, 413)

    def test_bigquery_failure_returns_a_clear_error_without_internals(self):
        token = encrypt_token({"kind": "download", "sql": SQL, "language": "en"})
        with (
            patch("app.validate_sql", side_effect=lambda sql, row_limit=100: sql),
            patch("app.stream_query_rows", side_effect=RuntimeError("secret project detail")),
        ):
            response = self.client.get(f"/exports/{token}")

        self.assertEqual(response.status_code, 500)
        self.assertEqual(response.json()["detail"], "I couldn't generate the Excel file. Please try again.")


if __name__ == "__main__":
    unittest.main()
