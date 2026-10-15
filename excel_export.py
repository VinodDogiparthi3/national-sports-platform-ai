import base64
import json
import math
import os
import re
import tempfile
from datetime import date, datetime, time, timezone
from decimal import Decimal
from hashlib import sha256

from cryptography.fernet import Fernet, InvalidToken
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ERROR_CODES, ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from config import EXPORT_MAX_ROWS, EXPORT_SECRET_KEY


XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Excel allows 1,048,576 rows per sheet; one is the header.
EXCEL_SHEET_MAX_DATA_ROWS = 1_048_575
EXCEL_CELL_MAX_CHARS = 32_767
# Excel stores 15 significant digits; larger integers are written as text to stay exact.
EXCEL_MAX_EXACT_INT = 10**15


class ExportTokenError(Exception):
    pass


class ExportTooLargeError(Exception):
    pass


# --------------------------------------------------
# Intent detection
# --------------------------------------------------

_ARABIC_DIACRITICS = re.compile(r"[\u064b-\u0652\u0640]")

_FILE_FORMAT = re.compile(
    r"\b(excel|xlsx|xls|spreadsheet|worksheet)\b|[إأا]كس[يى]?ل"
)
_EXPORT_VERB = re.compile(
    r"\b(export|download)\b|تصدير|صدّ?ر|تنزيل|نزّ?ل|تحميل|حمّ?ل"
)
_EXPORT_TARGET = re.compile(
    r"\b(data|results?|report|records?|table|analysis|list|rows|dataset|"
    r"everything|all|complete|full|this|these|above|it)\b"
    r"|البيانات|النتائج|التقرير|السجلات|الجدول|التحليل|هذه|هذا|الكاملة|كامل"
)


def is_excel_export_request(question: str) -> bool:
    text = _ARABIC_DIACRITICS.sub("", question.casefold())
    if _FILE_FORMAT.search(text):
        return True
    return bool(_EXPORT_VERB.search(text) and _EXPORT_TARGET.search(text))


# --------------------------------------------------
# Signed + encrypted tokens (stateless across Cloud Run instances)
# --------------------------------------------------

def _build_fernet() -> Fernet:
    secret = EXPORT_SECRET_KEY.encode() if EXPORT_SECRET_KEY else os.urandom(32)
    return Fernet(base64.urlsafe_b64encode(sha256(secret).digest()))


_fernet = _build_fernet()


def encrypt_token(payload: dict) -> str:
    return _fernet.encrypt(json.dumps(payload, ensure_ascii=False).encode()).decode()


def decrypt_token(token: str, ttl_seconds: int, kind: str) -> dict:
    try:
        payload = json.loads(_fernet.decrypt(token.encode(), ttl=ttl_seconds))
    except (InvalidToken, ValueError, UnicodeError) as error:
        raise ExportTokenError("Invalid or expired token.") from error

    if not isinstance(payload, dict) or payload.get("kind") != kind:
        raise ExportTokenError("Unexpected token type.")
    return payload


# --------------------------------------------------
# Workbook generation
# --------------------------------------------------

def friendly_header(name: str) -> str:
    if re.fullmatch(r"[a-z0-9_]+", name):
        return " ".join(part for part in name.split("_") if part).title() or name
    return name


def unique_headers(headers: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    result = []
    for header in headers:
        count = seen.get(header, 0) + 1
        seen[header] = count
        result.append(header if count == 1 else f"{header} ({count})")
    return result


def _plain_value(value):
    if value is None or isinstance(value, bool):
        return value

    if isinstance(value, int):
        return str(value) if abs(value) >= EXCEL_MAX_EXACT_INT else value

    if isinstance(value, float):
        return value if math.isfinite(value) else str(value)

    if isinstance(value, Decimal):
        return value if value.is_finite() else str(value)

    if isinstance(value, datetime):
        if value.tzinfo is not None:
            return value.astimezone(timezone.utc).replace(tzinfo=None)
        return value

    if isinstance(value, time):
        return value.replace(tzinfo=None)

    if isinstance(value, date):
        return value

    if isinstance(value, (bytes, bytearray)):
        return base64.b64encode(bytes(value)).decode()

    if isinstance(value, (list, tuple, dict)):
        return json.dumps(value, default=str, ensure_ascii=False)

    return str(value)


def _cell(sheet, value):
    value = _plain_value(value)
    if not isinstance(value, str):
        return value

    value = ILLEGAL_CHARACTERS_RE.sub("", value)[:EXCEL_CELL_MAX_CHARS]
    if value.startswith("=") or value in ERROR_CODES:
        # Keep user/database text as text so Excel never evaluates it as a formula.
        cell = WriteOnlyCell(sheet, value=value)
        cell.data_type = "s"
        return cell
    return value


def _abort_workbook(workbook) -> None:
    # Write-only sheets keep their own temp files until a successful save (private API, version pinned).
    for sheet in workbook.worksheets:
        try:
            sheet.close()
        except Exception:
            pass
        writer = getattr(sheet, "_writer", None)
        if writer is not None:
            try:
                writer.cleanup()
            except Exception:
                pass


def write_workbook(
    headers: list[str],
    rows,
    *,
    language: str = "en",
    max_rows: int = EXPORT_MAX_ROWS,
    sheet_max_rows: int = EXCEL_SHEET_MAX_DATA_ROWS,
) -> tuple[str, int]:
    """Stream rows into an .xlsx temp file; returns (path, row count). Caller deletes the file."""
    headers = unique_headers(headers)
    base_title = "البيانات" if language == "ar" else "Data"
    workbook = Workbook(write_only=True)
    state = {"sheet": None, "count": 0}

    def start_sheet():
        index = len(workbook.worksheets) + 1
        sheet = workbook.create_sheet(base_title if index == 1 else f"{base_title} {index}")
        sheet.sheet_view.rightToLeft = language == "ar"
        sheet.freeze_panes = "A2"
        for position, header in enumerate(headers, start=1):
            sheet.column_dimensions[get_column_letter(position)].width = max(12, min(len(header) + 6, 50))

        header_cells = []
        for header in headers:
            cell = WriteOnlyCell(sheet, value=ILLEGAL_CHARACTERS_RE.sub("", header))
            cell.font = Font(bold=True)
            cell.data_type = "s"
            header_cells.append(cell)
        sheet.append(header_cells)
        state["sheet"] = sheet
        state["count"] = 0

    handle, path = tempfile.mkstemp(suffix=".xlsx")
    os.close(handle)
    total = 0
    try:
        for row in rows:
            if total >= max_rows:
                raise ExportTooLargeError(f"Export exceeds {max_rows:,} rows.")
            if state["sheet"] is None or state["count"] >= sheet_max_rows:
                start_sheet()

            state["sheet"].append([_cell(state["sheet"], value) for value in row])
            state["count"] += 1
            total += 1

        if state["sheet"] is None:
            start_sheet()
        workbook.save(path)
    except BaseException:
        _abort_workbook(workbook)
        if os.path.exists(path):
            os.remove(path)
        raise

    return path, total


def write_workbook_sheets(
    sheets,
    *,
    language: str = "en",
    max_rows: int = EXPORT_MAX_ROWS,
) -> tuple[str, int]:
    """Write each view's rows to a distinct worksheet, enforcing one workbook row cap."""
    workbook = Workbook(write_only=True)
    used_titles = set()

    def create_sheet(title, headers):
        base = re.sub(r"[\\/*?:\[\]]", " ", title).strip() or "Data"
        base = base[:31]
        candidate = base
        suffix = 2
        while candidate.casefold() in used_titles:
            ending = f" {suffix}"
            candidate = f"{base[:31 - len(ending)]}{ending}"
            suffix += 1
        used_titles.add(candidate.casefold())

        sheet = workbook.create_sheet(candidate)
        sheet.sheet_view.rightToLeft = language == "ar"
        sheet.freeze_panes = "A2"
        headers = unique_headers(headers)
        for position, header in enumerate(headers, start=1):
            sheet.column_dimensions[get_column_letter(position)].width = max(12, min(len(header) + 6, 50))
        header_cells = []
        for header in headers:
            cell = WriteOnlyCell(sheet, value=ILLEGAL_CHARACTERS_RE.sub("", header))
            cell.font = Font(bold=True)
            cell.data_type = "s"
            header_cells.append(cell)
        sheet.append(header_cells)
        return sheet

    handle, path = tempfile.mkstemp(suffix=".xlsx")
    os.close(handle)
    total = 0
    try:
        for title, headers, rows in sheets:
            sheet = create_sheet(title, headers)
            for row in rows:
                if total >= max_rows:
                    raise ExportTooLargeError(f"Export exceeds {max_rows:,} rows.")
                sheet.append([_cell(sheet, value) for value in row])
                total += 1
        if not workbook.worksheets:
            create_sheet("Data", [])
        workbook.save(path)
    except BaseException:
        _abort_workbook(workbook)
        if os.path.exists(path):
            os.remove(path)
        raise
    return path, total
